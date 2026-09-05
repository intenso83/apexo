from __future__ import annotations

import csv
import json
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from apexo_ocr.constants import IMPORT_HEADERS, INCLUDE_NO, REVIEWED  # noqa: E402
from apexo_ocr.apexo_import import (  # noqa: E402
    ApexoImportError,
    dry_run_apexo,
    import_apexo,
    validate_apexo_url,
)
from apexo_ocr.csv_contract import load_reviewed_csv  # noqa: E402
from apexo_ocr.documents import iter_document_pages  # noqa: E402
from apexo_ocr.models import OcrRecord  # noqa: E402
from apexo_ocr.pipeline import _mark_duplicates  # noqa: E402
from apexo_ocr.security import PrivacyBoundaryError, validate_loopback_endpoint  # noqa: E402
from apexo_ocr.validation import normalize_engine_result, valid_afm  # noqa: E402
from apexo_ocr.workbook import (  # noqa: E402
    WorkbookValidationError,
    create_review_workbook,
    export_reviewed_csv,
)


class PrivacyBoundaryTests(unittest.TestCase):
    def test_loopback_endpoints_are_allowed(self) -> None:
        self.assertEqual(
            validate_loopback_endpoint("http://127.0.0.1:11434/"),
            "http://127.0.0.1:11434",
        )
        self.assertEqual(
            validate_loopback_endpoint("http://localhost:11434"),
            "http://localhost:11434",
        )

    def test_remote_endpoints_are_rejected(self) -> None:
        for endpoint in (
            "https://example.com:11434",
            "http://192.168.1.9:11434",
            "http://ollama.example:11434",
        ):
            with self.subTest(endpoint=endpoint), self.assertRaises(PrivacyBoundaryError):
                validate_loopback_endpoint(endpoint)


class ValidationTests(unittest.TestCase):
    def test_normalization_preserves_leading_zeroes(self) -> None:
        result = normalize_engine_result(
            {
                "values": {
                    "last_name": " TEST ",
                    "first_name": "PATIENT",
                    "postal_code": "05463",
                    "amka": "010101-00000",
                    "afm": "012345678",
                    "registration_date_raw": "02/01/2026",
                },
                "confidence": {"last_name": 0.9},
                "fields_need_review": [],
                "notes": [],
            },
            source_file="synthetic.png",
            source_page=1,
            source_sha256="a" * 64,
            engine="test",
        )
        self.assertEqual(result.values["postal_code"], "05463")
        self.assertEqual(result.values["amka"], "01010100000")
        self.assertEqual(result.values["afm"], "012345678")
        self.assertEqual(result.values["registration_date"], "2026-01-02")
        self.assertFalse(valid_afm("012345678"))

    def test_duplicate_is_visible_and_excluded_from_preview(self) -> None:
        first = _record(1, amka="01010100000")
        second = _record(2, amka="01010100000")
        _mark_duplicates([first, second])
        self.assertEqual(second.include_in_import, INCLUDE_NO)
        self.assertEqual(second.duplicate_of, first.source_label)
        self.assertIn("duplicate", second.fields_need_review)


class WorkbookTests(unittest.TestCase):
    def test_workbook_and_review_gate(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "review.xlsx"
            csv_path = root / "patients_import.csv"
            record = _record(1)
            record.values["address"] = '=HYPERLINK("https://invalid")'
            create_review_workbook(
                [record],
                workbook_path,
                batch_name="TEST_BATCH",
                requested_engine="test",
                selected_engine="test",
                model="",
            )

            workbook = load_workbook(workbook_path, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "README",
                    "OCR_Review",
                    "Import_Preview",
                    "Validation_Report",
                    "Registration_Dates",
                    "Lists",
                ],
            )
            review = workbook["OCR_Review"]
            review_headers = [review.cell(1, column).value for column in range(1, review.max_column + 1)]
            address = review.cell(2, review_headers.index("address") + 1)
            self.assertEqual(address.data_type, "s")
            self.assertTrue(str(address.value).startswith("="))
            preview = workbook["Import_Preview"]
            preview_headers = [preview.cell(1, column).value for column in range(1, preview.max_column + 1)]
            self.assertEqual(preview_headers, IMPORT_HEADERS)
            for name in ("history_reason", "history_medicines"):
                self.assertIn(preview.cell(2, preview_headers.index(name) + 1).value, (None, ""))
            for name in ("hist_hypertension", "hist_coagulation"):
                self.assertEqual(preview.cell(2, preview_headers.index(name) + 1).value, "0")
            workbook.close()

            with self.assertRaises(WorkbookValidationError):
                export_reviewed_csv(workbook_path, csv_path)

            workbook = load_workbook(workbook_path)
            review = workbook["OCR_Review"]
            review_headers = [review.cell(1, column).value for column in range(1, review.max_column + 1)]
            review.cell(2, review_headers.index("review_status") + 1).value = REVIEWED
            workbook.save(workbook_path)
            workbook.close()
            export_reviewed_csv(workbook_path, csv_path)
            with csv_path.open("r", encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.reader(stream))
            self.assertEqual(rows[0], IMPORT_HEADERS)
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[1][IMPORT_HEADERS.index("index")], "0")

    def test_contact_only_export_rejects_medical_values(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "review.xlsx"
            csv_path = root / "patients_import.csv"
            create_review_workbook(
                [_record(1)],
                workbook_path,
                batch_name="TEST_BATCH",
                requested_engine="test",
                selected_engine="test",
                model="",
            )
            workbook = load_workbook(workbook_path)
            review = workbook["OCR_Review"]
            review_headers = [review.cell(1, column).value for column in range(1, review.max_column + 1)]
            review.cell(2, review_headers.index("review_status") + 1).value = REVIEWED
            preview = workbook["Import_Preview"]
            preview_headers = [preview.cell(1, column).value for column in range(1, preview.max_column + 1)]
            preview.cell(2, preview_headers.index("history_medicines") + 1).value = "not allowed"
            workbook.save(workbook_path)
            workbook.close()
            with self.assertRaises(WorkbookValidationError):
                export_reviewed_csv(workbook_path, csv_path)


class PdfRenderingTests(unittest.TestCase):
    def test_scanned_pdf_is_rendered_without_external_tools(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic.pdf"
            image = Image.new("RGB", (600, 800), "white")
            image.save(path, "PDF", resolution=150)
            pages = list(iter_document_pages(path, dpi=150))
            self.assertEqual(len(pages), 1)
            self.assertGreater(pages[0].image.width, 500)
            pages[0].image.close()


class ReviewedCsvTests(unittest.TestCase):
    def test_valid_contact_csv_and_medical_scope_gate(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "reviewed.csv"
            _write_reviewed_csv(path, [_csv_values(index="0")])
            report = load_reviewed_csv(path)
            self.assertTrue(report.valid)
            self.assertEqual(len(report.rows), 1)
            self.assertEqual(report.rows[0].values["postal_code"], "05463")

            values = _csv_values(index="0")
            values["history_medicines"] = "not allowed"
            _write_reviewed_csv(path, [values])
            report = load_reviewed_csv(path)
            self.assertFalse(report.valid)
            self.assertIn("contact_only_scope", {issue.code for issue in report.errors})

    def test_duplicate_rows_are_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "reviewed.csv"
            first = _csv_values(index="0")
            second = _csv_values(index="1")
            _write_reviewed_csv(path, [first, second])
            report = load_reviewed_csv(path)
            self.assertFalse(report.valid)
            self.assertEqual(
                sum(issue.code == "duplicate_inside_csv" for issue in report.errors), 2
            )


class ApexoImportTests(unittest.TestCase):
    def test_url_policy(self) -> None:
        self.assertEqual(validate_apexo_url("https://example.test/"), "https://example.test")
        self.assertEqual(validate_apexo_url("http://127.0.0.1:8090"), "http://127.0.0.1:8090")
        with self.assertRaises(ApexoImportError):
            validate_apexo_url("http://192.168.1.10:8090")

    def test_dry_run_import_and_idempotent_repeat(self) -> None:
        with tempfile.TemporaryDirectory() as directory, _FakeApexoServer() as server:
            path = Path(directory) / "reviewed.csv"
            _write_reviewed_csv(path, [_csv_values(index="0")])
            report = load_reviewed_csv(path)
            plan = dry_run_apexo(report, server.url, "user@example.test", "secret")
            self.assertTrue(plan.can_import)
            self.assertEqual(plan.result["importable_count"], 1)
            result = import_apexo(report, plan, "user@example.test", "secret")
            self.assertEqual(result["created_count"], 1)
            self.assertEqual(len(server.records), 1)
            data = server.records[0]["data"]
            self.assertEqual(data["surname"], "SYNTHETIC")
            self.assertEqual(data["first_name"], "PATIENT")
            self.assertEqual(data["birth_date"], "1980-06-15")
            self.assertNotIn("history_medicines", data)

            repeat = dry_run_apexo(report, server.url, "user@example.test", "secret")
            self.assertFalse(repeat.can_import)
            self.assertEqual(repeat.result["already_imported_count"], 1)


def _record(page: int, *, amka: str = "") -> OcrRecord:
    return OcrRecord(
        source_file="synthetic.pdf",
        source_page=page,
        source_sha256="b" * 64,
        engine="test",
        values={
            "last_name": "SYNTHETIC",
            "first_name": f"PATIENT{page}",
            "amka": amka,
            "registration_date_raw": "01/01/2026",
            "registration_date": "2026-01-01",
        },
    )


def _csv_values(*, index: str) -> dict[str, str]:
    values = {name: "" for name in IMPORT_HEADERS}
    values.update(
        {
            "index": index,
            "last_name": "SYNTHETIC",
            "first_name": "PATIENT",
            "profession": "TESTER",
            "address": "1 TEST STREET",
            "city": "ATHENS",
            "area": "ATTICA",
            "postal_code": "05463",
            "birth_date": "1980-06-15",
            "registration_date": "2026-09-01",
            "mobile": "0691234567",
            "email": "synthetic@example.test",
        }
    )
    for name in IMPORT_HEADERS:
        if name.startswith("hist_"):
            values[name] = "0"
    return values


def _write_reviewed_csv(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=IMPORT_HEADERS, lineterminator="\r\n")
        writer.writeheader()
        writer.writerows(rows)


class _FakeApexoServer:
    def __init__(self) -> None:
        self.records: list[dict[str, object]] = []
        owner = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, format: str, *args: object) -> None:
                return

            def _json(self, status: int, value: dict[str, object]) -> None:
                payload = json.dumps(value).encode("utf-8")
                self.send_response(status)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)

            def do_POST(self) -> None:
                path = urlsplit(self.path).path
                length = int(self.headers.get("Content-Length", "0"))
                body = json.loads(self.rfile.read(length) or b"{}")
                if path.endswith("/_superusers/auth-with-password"):
                    self._json(404, {"message": "missing"})
                    return
                if path.endswith("/users/auth-with-password"):
                    if body.get("identity") == "user@example.test" and body.get("password") == "secret":
                        self._json(200, {"token": "test-token"})
                    else:
                        self._json(401, {"message": "denied"})
                    return
                if path == "/api/collections/data/records":
                    record = dict(body)
                    record["updated"] = f"2026-09-01 00:00:{len(owner.records):02d}.000Z"
                    owner.records.append(record)
                    self._json(200, record)
                    return
                self._json(404, {"message": "missing"})

            def do_GET(self) -> None:
                path = urlsplit(self.path).path
                if path == "/api/collections/data/records":
                    self._json(
                        200,
                        {
                            "page": 1,
                            "perPage": 200,
                            "totalItems": len(owner.records),
                            "totalPages": 1,
                            "items": owner.records,
                        },
                    )
                    return
                self._json(404, {"message": "missing"})

            def do_DELETE(self) -> None:
                path = urlsplit(self.path).path
                record_id = path.rsplit("/", 1)[-1]
                owner.records[:] = [record for record in owner.records if record.get("id") != record_id]
                self.send_response(204)
                self.end_headers()

        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.url = f"http://127.0.0.1:{self.httpd.server_address[1]}"

    def __enter__(self) -> "_FakeApexoServer":
        self.thread.start()
        return self

    def __exit__(self, exc_type: object, exc: object, traceback: object) -> None:
        self.httpd.shutdown()
        self.httpd.server_close()
        self.thread.join(timeout=5)


if __name__ == "__main__":
    unittest.main()
