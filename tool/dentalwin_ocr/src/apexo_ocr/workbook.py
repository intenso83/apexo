from __future__ import annotations

import csv
import os
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from .constants import (
    HISTORY_FLAG_HEADERS,
    HISTORY_TEXT_HEADERS,
    IMPORT_HEADERS,
    INCLUDE_YES,
    OCR_REVIEW_HEADERS,
    REGISTRATION_DATE_HEADERS,
    REVIEWED,
    REVIEW_REQUIRED,
    VALIDATION_HEADERS,
)
from .models import OcrRecord


class WorkbookValidationError(ValueError):
    pass


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBTITLE_FILL = PatternFill("solid", fgColor="D9EAF7")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="FCE4D6")
GOOD_FILL = PatternFill("solid", fgColor="E2F0D9")


def create_review_workbook(
    records: list[OcrRecord],
    output_path: str | Path,
    *,
    batch_name: str,
    requested_engine: str,
    selected_engine: str,
    model: str,
) -> Path:
    path = Path(output_path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    review = workbook.create_sheet("OCR_Review")
    preview = workbook.create_sheet("Import_Preview")
    validation = workbook.create_sheet("Validation_Report")
    registrations = workbook.create_sheet("Registration_Dates")
    lists = workbook.create_sheet("Lists")

    _populate_readme(
        readme,
        records,
        batch_name=batch_name,
        requested_engine=requested_engine,
        selected_engine=selected_engine,
        model=model,
    )
    _populate_review(review, records)
    _populate_preview(preview, records)
    _populate_validation(validation, records)
    _populate_registration_dates(registrations, records)
    _populate_lists(lists)
    lists.sheet_state = "veryHidden"

    workbook.properties.title = f"{batch_name} OCR review"
    workbook.properties.subject = "Local contact-only OCR review for DentalWin import"
    workbook.properties.creator = "Apexo OCR Utility"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    _atomic_save(workbook, path)
    return path


def _populate_readme(
    sheet: Any,
    records: list[OcrRecord],
    *,
    batch_name: str,
    requested_engine: str,
    selected_engine: str,
    model: str,
) -> None:
    rows = [
        ["Apexo OCR Utility — contact-only review workbook", ""],
        ["Batch", batch_name],
        ["Forms/pages", str(len(records))],
        ["Requested engine", requested_engine],
        ["Engine used", selected_engine],
        ["Local model", model if selected_engine == "ollama" else ""],
        ["Privacy", "Local processing only. No remote OCR endpoint is accepted."],
        ["Scope", "Contact/administrative fields + registration_date only."],
        [
            "Excluded",
            "Medical history, diagnoses, medicines, health checkboxes, signatures, treatment data.",
        ],
        ["", ""],
        ["SAFE WORKFLOW", ""],
        ["1", "Review OCR_Review and compare every value with the original scan."],
        ["2", "Correct the final values in Import_Preview. Keep identifiers as text."],
        [
            "3",
            f"For every included page set OCR_Review.review_status to {REVIEWED} only after checking it.",
        ],
        [
            "4",
            "Use Apexo OCR Utility → Export reviewed CSV. Do not use Excel Save As CSV.",
        ],
        ["5", "Run the DentalWin dry run before any copy/live import."],
        ["", ""],
        [
            "Important",
            "All OCR rows start as ΧΡΕΙΑΖΕΤΑΙ ΕΛΕΓΧΟ. The utility blocks CSV export until review is complete.",
        ],
        [
            "Duplicates",
            "A detected duplicate is visible in OCR_Review and excluded from Import_Preview. Confirm manually.",
        ],
    ]
    for row in rows:
        sheet.append(row)
    sheet.merge_cells("A1:B1")
    sheet["A1"].fill = HEADER_FILL
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    sheet["A1"].alignment = Alignment(horizontal="center")
    for row in (11,):
        sheet[f"A{row}"].fill = SUBTITLE_FILL
        sheet[f"A{row}"].font = Font(bold=True)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    for row in range(1, sheet.max_row + 1):
        sheet[f"A{row}"].font = Font(bold=True) if sheet[f"A{row}"].value else Font()
        sheet[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 100
    sheet.freeze_panes = "A2"


def _populate_review(sheet: Any, records: list[OcrRecord]) -> None:
    _append_text_row(sheet, OCR_REVIEW_HEADERS)
    for record in records:
        row = record.review_row()
        _append_text_row(sheet, [row.get(header, "") for header in OCR_REVIEW_HEADERS])
    _format_data_sheet(sheet, OCR_REVIEW_HEADERS, "OCRReviewTable")
    if records:
        include_col = OCR_REVIEW_HEADERS.index("include_in_import") + 1
        status_col = OCR_REVIEW_HEADERS.index("review_status") + 1
        include_validation = DataValidation(type="list", formula1='"YES,NO"', allow_blank=False)
        status_validation = DataValidation(
            type="list", formula1=f'"{REVIEW_REQUIRED},{REVIEWED}"', allow_blank=False
        )
        sheet.add_data_validation(include_validation)
        sheet.add_data_validation(status_validation)
        include_validation.add(
            f"{get_column_letter(include_col)}2:{get_column_letter(include_col)}{sheet.max_row}"
        )
        status_validation.add(
            f"{get_column_letter(status_col)}2:{get_column_letter(status_col)}{sheet.max_row}"
        )
        status_letter = get_column_letter(status_col)
        status_range = f"{status_letter}2:{status_letter}{sheet.max_row}"
        sheet.conditional_formatting.add(
            status_range,
            FormulaRule(formula=[f'${status_letter}2="{REVIEW_REQUIRED}"'], fill=WARNING_FILL),
        )
        sheet.conditional_formatting.add(
            status_range,
            FormulaRule(formula=[f'${status_letter}2="{REVIEWED}"'], fill=GOOD_FILL),
        )
    widths = {
        "source_file": 28,
        "source_page": 11,
        "source_sha256": 18,
        "include_in_import": 16,
        "duplicate_of": 24,
        "review_status": 24,
        "engine": 12,
        "model": 30,
        "full_name_raw": 30,
        "last_name": 20,
        "first_name": 20,
        "profession": 22,
        "address": 32,
        "city": 18,
        "area": 18,
        "postal_code": 13,
        "phone": 18,
        "mobile": 18,
        "email": 32,
        "birth_date_raw": 18,
        "birth_date": 16,
        "amka": 17,
        "afm": 14,
        "doy": 18,
        "referrer_raw": 36,
        "referrer": 28,
        "registration_date_raw": 22,
        "registration_date": 20,
        "average_confidence": 18,
        "fields_need_review": 40,
        "extraction_notes": 55,
        "raw_ocr_contact_zones": 80,
    }
    _set_widths(sheet, OCR_REVIEW_HEADERS, widths)


def _populate_preview(sheet: Any, records: list[OcrRecord]) -> None:
    _append_text_row(sheet, IMPORT_HEADERS)
    included = [record for record in records if record.include_in_import == INCLUDE_YES]
    for index, record in enumerate(included):
        row = record.import_row(index)
        _append_text_row(sheet, [row.get(header, "") for header in IMPORT_HEADERS])
    _format_data_sheet(sheet, IMPORT_HEADERS, "ImportPreviewTable")
    widths = {header: 17 for header in IMPORT_HEADERS}
    widths.update(
        {
            "index": 8,
            "last_name": 20,
            "first_name": 20,
            "profession": 22,
            "address": 32,
            "city": 18,
            "area": 18,
            "postal_code": 13,
            "doy": 18,
            "email": 32,
            "referrer": 28,
            "patient_notes": 25,
        }
    )
    _set_widths(sheet, IMPORT_HEADERS, widths)
    if sheet.max_row >= 2:
        last_name = get_column_letter(IMPORT_HEADERS.index("last_name") + 1)
        first_name = get_column_letter(IMPORT_HEADERS.index("first_name") + 1)
        row_range = f"A2:{get_column_letter(len(IMPORT_HEADERS))}{sheet.max_row}"
        sheet.conditional_formatting.add(
            row_range,
            FormulaRule(
                formula=[f'OR(${last_name}2="",${first_name}2="")'],
                fill=ERROR_FILL,
            ),
        )


def _populate_validation(sheet: Any, records: list[OcrRecord]) -> None:
    _append_text_row(sheet, VALIDATION_HEADERS)
    for record in records:
        for issue in record.validation_issues:
            _append_text_row(
                sheet,
                [
                    record.source_file,
                    str(record.source_page),
                    issue.severity,
                    issue.field,
                    issue.code,
                    issue.message,
                ],
            )
        if record.duplicate_of:
            _append_text_row(
                sheet,
                [
                    record.source_file,
                    str(record.source_page),
                    "warning",
                    "duplicate",
                    "duplicate_candidate",
                    f"Πιθανό duplicate του {record.duplicate_of}; αποκλείστηκε από το preview.",
                ],
            )
    _format_data_sheet(sheet, VALIDATION_HEADERS, "ValidationReportTable")
    _set_widths(
        sheet,
        VALIDATION_HEADERS,
        {"source_file": 28, "source_page": 11, "severity": 12, "field": 24, "code": 28, "message": 65},
    )


def _populate_registration_dates(sheet: Any, records: list[OcrRecord]) -> None:
    _append_text_row(sheet, REGISTRATION_DATE_HEADERS)
    for record in records:
        _append_text_row(
            sheet,
            [
                record.source_file,
                str(record.source_page),
                record.values.get("registration_date_raw", ""),
                record.values.get("registration_date", ""),
                f"{record.confidence.get('registration_date', 0.0):.2f}",
                record.review_status,
            ],
        )
    _format_data_sheet(sheet, REGISTRATION_DATE_HEADERS, "RegistrationDatesTable")
    _set_widths(
        sheet,
        REGISTRATION_DATE_HEADERS,
        {
            "source_file": 28,
            "source_page": 11,
            "registration_date_raw": 24,
            "registration_date": 20,
            "confidence": 14,
            "review_status": 24,
        },
    )


def _populate_lists(sheet: Any) -> None:
    for row in (("include_in_import", "review_status"), ("YES", REVIEW_REQUIRED), ("NO", REVIEWED)):
        _append_text_row(sheet, row)


def _append_text_row(sheet: Any, values: Iterable[Any]) -> None:
    sheet.append(["" if value is None else str(value) for value in values])
    for cell in sheet[sheet.max_row]:
        if cell.value is not None:
            cell.data_type = "s"


def _format_data_sheet(sheet: Any, headers: list[str], table_name: str) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if sheet.max_row >= 2:
        table = Table(
            displayName=table_name,
            ref=f"A1:{get_column_letter(len(headers))}{sheet.max_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)


def _set_widths(sheet: Any, headers: list[str], widths: dict[str, float]) -> None:
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(header, 18)


def _atomic_save(workbook: Workbook, path: Path) -> None:
    descriptor, temp_name = tempfile.mkstemp(prefix=f".{path.stem}_", suffix=".xlsx", dir=path.parent)
    os.close(descriptor)
    temp_path = Path(temp_name)
    try:
        workbook.save(temp_path)
        temp_path.replace(path)
    finally:
        temp_path.unlink(missing_ok=True)


def export_reviewed_csv(
    workbook_path: str | Path,
    csv_path: str | Path,
    *,
    allow_unreviewed: bool = False,
) -> Path:
    source = Path(workbook_path).expanduser().resolve()
    target = Path(csv_path).expanduser().resolve()
    if source == target:
        raise WorkbookValidationError("Το CSV output δεν μπορεί να είναι το ίδιο αρχείο με το workbook.")
    try:
        workbook = load_workbook(source, data_only=False, read_only=False)
    except Exception as exc:
        raise WorkbookValidationError("Δεν άνοιξε το review workbook.") from exc
    try:
        if "OCR_Review" not in workbook.sheetnames or "Import_Preview" not in workbook.sheetnames:
            raise WorkbookValidationError("Λείπουν τα απαιτούμενα φύλλα OCR_Review / Import_Preview.")
        review = workbook["OCR_Review"]
        preview = workbook["Import_Preview"]
        review_headers = _read_headers(review)
        for required in ("include_in_import", "review_status"):
            if required not in review_headers:
                raise WorkbookValidationError(f"Λείπει η στήλη {required} από το OCR_Review.")
        include_index = review_headers.index("include_in_import") + 1
        status_index = review_headers.index("review_status") + 1
        included_review_rows = 0
        pending_rows: list[int] = []
        for row_number in range(2, review.max_row + 1):
            include = _cell_text(review.cell(row_number, include_index).value).upper()
            if include != INCLUDE_YES:
                continue
            included_review_rows += 1
            status = _cell_text(review.cell(row_number, status_index).value).upper()
            if status not in {REVIEWED, "REVIEWED"}:
                pending_rows.append(row_number)
        if pending_rows and not allow_unreviewed:
            sample = ", ".join(str(number) for number in pending_rows[:12])
            suffix = "…" if len(pending_rows) > 12 else ""
            raise WorkbookValidationError(
                f"Υπάρχουν {len(pending_rows)} included γραμμές χωρίς ολοκληρωμένο review "
                f"(OCR_Review rows: {sample}{suffix})."
            )

        headers = _read_headers(preview)
        if headers != IMPORT_HEADERS:
            raise WorkbookValidationError("Τα headers/order του Import_Preview δεν ταιριάζουν με το V6/V7 schema.")
        rows: list[list[str]] = []
        for row_number in range(2, preview.max_row + 1):
            row = [_cell_text(preview.cell(row_number, column).value) for column in range(1, len(headers) + 1)]
            if not any(row):
                continue
            if not row[IMPORT_HEADERS.index("last_name")] or not row[IMPORT_HEADERS.index("first_name")]:
                raise WorkbookValidationError(
                    f"Import_Preview row {row_number}: first_name/last_name είναι υποχρεωτικά."
                )
            for name in HISTORY_TEXT_HEADERS:
                if row[IMPORT_HEADERS.index(name)]:
                    raise WorkbookValidationError(
                        f"Import_Preview row {row_number}: το contact-only export δεν επιτρέπει {name}."
                    )
            for name in HISTORY_FLAG_HEADERS:
                value = row[IMPORT_HEADERS.index(name)]
                if value not in ("", "0", "0.0", "FALSE", "False", "false"):
                    raise WorkbookValidationError(
                        f"Import_Preview row {row_number}: το contact-only export δεν επιτρέπει {name}."
                    )
                row[IMPORT_HEADERS.index(name)] = "0"
            row[IMPORT_HEADERS.index("patient_notes")] = ""
            row[IMPORT_HEADERS.index("index")] = str(len(rows))
            rows.append(row)
        if not rows:
            raise WorkbookValidationError("Το Import_Preview δεν περιέχει γραμμές για export.")
        if included_review_rows != len(rows):
            raise WorkbookValidationError(
                "Ο αριθμός included rows στο OCR_Review δεν συμφωνεί με το Import_Preview. "
                "Διόρθωσε το workbook ή ξαναδημιούργησέ το."
            )
    finally:
        workbook.close()

    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temp_name = tempfile.mkstemp(prefix=f".{target.stem}_", suffix=".csv", dir=target.parent)
    os.close(descriptor)
    temp_path = Path(temp_name)
    try:
        with temp_path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, lineterminator="\r\n")
            writer.writerow(IMPORT_HEADERS)
            writer.writerows(rows)
        temp_path.replace(target)
    finally:
        temp_path.unlink(missing_ok=True)
    return target


def _read_headers(sheet: Any) -> list[str]:
    return [_cell_text(sheet.cell(1, column).value) for column in range(1, sheet.max_column + 1)]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
